# -*- coding: utf-8 -*-
"""
ops/import_survey.py —— 导入真实报名表（**选项 A：去标识粗粒度**）

用户决定：**敏感信息不记录，按选项 A 进行。**

所以本导入器**绝不写入**：
  · 称呼/姓名（第 2 列）—— 里面有真名
  · 联系方式（第 5 列）—— 微信号/手机
  · 「学校+科室+角色」原文 —— "中国中医科学院广安门医院/内窥镜中心（脾胃科）/2027届硕士毕业生"
    这种原文**本身就是指纹**，逐字存下来等于没去标识

只写入**粗粒度、非个人**的维度，且**尽量复用已有字段**（本库纪律：一个口径一份实现，不另造并行结构）：

| 报名表内容 | 落到哪 | 说明 |
|---|---|---|
| 关注方向（"药企/AI/金融"…） | `F_PSN_INT_INTEREST_DOMAIN`（码表 `CT_INTEREST_DOMAIN`，IN01–IN99） | **已有字段** |
| 学历阶段（"研一""博士"…） | `F_PSN_EDU_DEGREE_LEVEL`（`CT_DEGREE_LEVEL`） | **已有字段**；只在**明确**时映射 |
| 是否临床医学类 | `F_PSN_EDU_IS_CLINICAL_MAJOR`（`CT_YES_NO`） | **已有字段**；只在**明确**时映射 |
| 参与状态（"倾听"/"可分享经验"） | `F_PSN_SURVEY_PARTICIPATION` | **新增一个字段**（字典里确实没有） |

⚠ 两条纪律：
  ① **映射不了的绝不猜** —— 记为该字段"没有值"（不写行），并**计数上报**。
     本库原则：缺失 ≠ 空值；未映射要计数，不要静默丢弃。
  ② 缺失本身不是问题。报名表本来就只有 5 个字段，**没有履历、没有岗位**是正常的，
     问题只会出在"把缺失当成 0"。

用法：
  python ops/import_survey.py --plan      # 只解析与报告，不写库
  python ops/import_survey.py --apply     # 写入
  python ops/import_survey.py --cleanup   # 撤销本次导入
"""
from __future__ import annotations

import argparse
import io
import os
import re
import sys
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE, "ops"))

from openpyxl import load_workbook                       # noqa: E402
import psycopg                                           # noqa: E402
from psycopg.rows import dict_row                        # noqa: E402

DSN = ("host=127.0.0.1 port=55432 dbname=medtalent user=postgres connect_timeout=10 "
       "options='-c search_path=mt,public'")

SRC_DIR = (r"D:\HuaweiMoveData\Users\团聚体\Documents\xwechat_files"
           r"\wxid_t9wwitxtv4xo22_8114\msg\file\2026-10")
FILES = ["医疗去向探索cc（北京场总人数(1).xlsx", "医疗去向探索cc(1).xlsx",
         "医学生就业方向与转行路径分析王炸(1).xlsx", "医学生职业发展交流表无糖(2).xlsx"]

SRC_ID = "src_survey_2026_10"
SRC_SYS = "survey2026_10"
PREFIX = "per_svy26_"          # 本次导入的 person_id 前缀（cleanup 用）
BATCH_ATTR = "survey_batch"

# ---- 关注方向：关键词 → CT_INTEREST_DOMAIN 的码 ----
# **只写有明确对应的**；没有对应的词落进"未映射"并计数，不硬塞。
KW2IN = [
    (("药企", "药物", "医药", "药品", "注册"), ["IN02"]),
    (("器械", "IVD", "ivd", "设备"), ["IN03"]),
    (("CRO", "cro", "临床试验"), ["IN04"]),
    (("生物技术", "基因"), ["IN05"]),
    (("AI", "ai", "人工智能", "算法"), ["IN06"]),
    (("数据", "统计", "建模"), ["IN07"]),
    (("互联网", "数字健康", "数字医疗"), ["IN08"]),
    (("公卫", "流行病"), ["IN09"]),
    (("保险", "支付"), ["IN10"]),
    (("医院管理", "运营"), ["IN11"]),
    (("咨询",), ["IN12"]),
    (("投资", "证券", "PE", "pe", "一级市场", "股权投资"), ["IN13"]),
    (("金融", "风险"), ["IN14"]),
    (("法务", "知识产权", "专利"), ["IN15"]),
    (("科研", "学术", "读博", "读研", "实验室"), ["IN16"]),
    (("出版", "编辑", "科普"), ["IN17"]),
    (("教育", "培训", "教师"), ["IN18"]),
    (("市场", "品牌", "公关"), ["IN19"]),
    (("销售", "商务", "BD"), ["IN20"]),
    (("产品", "用户研究"), ["IN21"]),
    (("供应链", "生产", "制造"), ["IN22"]),
    (("政府", "体制", "公务员", "选调"), ["IN23"]),
    (("创业", "自由职业"), ["IN24"]),
    (("软件", "开发", "程序员"), ["IN26"]),
    (("临床", "中医", "医生", "医疗健康"), ["IN01"]),
]
# 明确无对应码表的通用词（**不是**关注方向，单独计数，不塞进 IN99）
IGNORE_WORDS = ["转行", "未来", "发展", "方向", "关注", "问题", "探索", "交流",
                "好奇", "了解", "分享", "经验", "生活", "感受", "规划", "选择",
                "路径", "卡点", "跨界", "切入", "成长", "机会"]


def degree_of(t: str):
    if re.search(r"八年制|本博连读", t):
        return "D6"
    if re.search(r"博士后", t):
        return "D5"
    if re.search(r"博士", t):
        return "D4"
    if re.search(r"研一|研二|研三|硕士|研究生在读|MBA|mba", t):
        return "D3"
    if re.search(r"本科|大[一二三四五]|应届", t):
        return "D2"
    if re.search(r"专科|高职", t):
        return "D1"
    return None


def clinical_of(t: str):
    if re.search(r"临床|中医|口腔|护理|儿科|眼科学|重症|心血管|泌外|影像|麻醉|肿瘤|皮肤|精神|妇产", t):
        return "Y"
    return None


def interests_of(t: str):
    codes, unmapped = [], []
    for kws, cs in KW2IN:
        if any(k in t for k in kws):
            for c in cs:
                if c not in codes:
                    codes.append(c)
    # 未映射：把文本切成片段，去掉能映射的与无意义的，剩下的才是"真的没映射"
    if not codes:
        seg = re.split(r"[\s,，、;；/。.（）()【】\[\]+&\-—:：!！?？]+", t)
        for s in seg:
            s = s.strip()
            if len(s) < 2:
                continue
            if any(w in s for w in IGNORE_WORDS):
                continue
            if any(any(k in s for k in kws) for kws, _ in KW2IN):
                continue
            if re.fullmatch(r"[\dA-Za-z]{1,3}", s):
                continue
            unmapped.append(s)
    return sorted(codes), unmapped[:6]


def parse():
    """解析 4 份文件 → 去标识的记录列表（原称呼/联系方式**读入即丢**，不进返回值）。"""
    seen, recs = set(), []
    for fn in FILES:
        p = os.path.join(SRC_DIR, fn)
        if not os.path.isfile(p):
            print("[!] 缺文件：%s" % fn)
            continue
        wb = load_workbook(p, read_only=True, data_only=True)
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
        for r in rows:
            c = lambda i: "" if i >= len(r) or r[i] is None else str(r[i]).strip()
            if "怎么称呼您" in c(1) or not c(1):
                continue
            # 去重键：用**原文**在内存里判重，但原文不进库
            key = re.sub(r"\s+", "", c(2))
            if key in seen:
                continue
            seen.add(key)
            topic = c(3) + " " + c(7)
            codes, unmapped = interests_of(topic)
            recs.append({
                "org_text_hash": __import__("hashlib").sha1(key.encode()).hexdigest()[:10],
                "degree": degree_of(c(2)),
                "clinical": clinical_of(c(2)),
                "interests": codes,
                "unmapped": unmapped,
                "state": ("可分享经验" if "分享" in c(4) else "倾听"),
            })
    return recs


def ensure_schema(c):
    """登记来源 + 新增字段（幂等）。"""
    # source_type 用码值：SRC5 = 问卷（码表 CT_SOURCE_TYPE）
    c.execute("""INSERT INTO mt.source_registry
                     (source_id, name, source_type, is_synthetic, license_note)
                 VALUES (%s, %s, 'SRC5', false, %s)
                 ON CONFLICT (source_id) DO UPDATE
                   SET name = EXCLUDED.name, source_type = 'SRC5',
                       is_synthetic = false, license_note = EXCLUDED.license_note""",
              (SRC_ID, "医学生去向交流会报名表（2026-10，去标识）",
               "**真实数据**。来源是交流会报名表；按用户决定采用**去标识**导入："
               "不记录称呼、联系方式与「学校+科室+角色」原文，只保留粗粒度维度"
               "（关注方向 / 学历阶段 / 是否临床 / 参与状态）。"
               "⚠ 报名≠同意进入人才库，所以本来源**不含任何个人标识信息**。"))
    c.execute("""INSERT INTO mt.field_catalog
                   (field_id, entity_id, title, description, data_type, value_type,
                    cardinality, is_arrayed, collection_method, is_private,
                    access_tier, status, version, is_required, form_section)
                 VALUES ('F_PSN_SURVEY_PARTICIPATION','person','交流活动参与状态',
                         '报名表里的「倾向活动状态」：倾听 / 可分享经验。'
                         '**去标识**：与姓名联系方式无关，只表示参与意愿。',
                         'string','string',1,false,'survey',false,'T1','active',1,false,'活动')
                 ON CONFLICT (field_id) DO NOTHING""")
    c.execute("""INSERT INTO mt.attribute_definition
                   (attr_key, entity_id, title, data_type, status, promoted_to_field)
                 VALUES ('survey_participation','person','交流活动参与状态','text','active',
                         'F_PSN_SURVEY_PARTICIPATION')
                 ON CONFLICT (attr_key) DO NOTHING""")


def apply_(recs):
    with psycopg.connect(DSN, row_factory=dict_row) as c:
        ensure_schema(c)
        n = 0
        for i, r in enumerate(recs, 1):
            pid = "%s%03d" % (PREFIX, i)
            c.execute("""INSERT INTO mt.person
                           (person_id, subject_code, access_tier, quality_flags,
                            confidence, verify_status, enroll_channel, source_id)
                         VALUES (%s,%s,'T1',ARRAY['real_survey_response'],0.9,'V1','survey',%s)
                         ON CONFLICT (person_id) DO UPDATE
                           SET quality_flags = EXCLUDED.quality_flags,
                               subject_code = EXCLUDED.subject_code,
                               source_id = EXCLUDED.source_id""",
                      (pid, "SVY26-%03d" % i, SRC_ID))
            n += 1
            def fv(field, code=None, codes=None, text=None):
                c.execute("""INSERT INTO mt.field_value
                               (value_id, field_id, subject_type, subject_id,
                                value_code, value_codes, value_text, source_id, evidence_grade)
                             VALUES (%s,%s,'person',%s,%s,%s,%s,%s,'B')
                             ON CONFLICT (value_id) DO NOTHING""",
                          ("fv_%s_%s" % (pid, field.split("_")[-1][:6]), field,
                           pid, code, codes, text, SRC_ID))
            if r["interests"]:
                fv("F_PSN_INT_INTEREST_DOMAIN", codes=r["interests"])
            if r["degree"]:
                fv("F_PSN_EDU_DEGREE_LEVEL", code=r["degree"])
            if r["clinical"]:
                fv("F_PSN_EDU_IS_CLINICAL_MAJOR", code=r["clinical"])
            fv("F_PSN_SURVEY_PARTICIPATION", text=r["state"])
            # ⚠ 未映射的片段**不写库** —— 它们里面是报名者的原话/问句
            #（如有没有儿科背景成功跑路的真实案例），逐条存下来就等于把
            # 自由文本带进了库，与选项 A 相悖。**只做聚合计数**（见 plan() 的输出）。
        c.commit()
    print("已导入 %d 人（去标识）" % n)
    return n


def cleanup():
    with psycopg.connect(DSN, row_factory=dict_row) as c:
        c.execute("DELETE FROM mt.field_value WHERE subject_id LIKE %s", (PREFIX + "%",))
        c.execute("DELETE FROM mt.external_identity WHERE person_id LIKE %s", (PREFIX + "%",))
        c.execute("DELETE FROM mt.person WHERE person_id LIKE %s", (PREFIX + "%",))
        c.commit()
    print("已撤销本次导入")


def plan(recs):
    tot = len(recs)
    print("=" * 88)
    print("选项 A（去标识）导入计划 —— 只解析，不写库")
    print("=" * 88)
    print("去重后 **%d 位**（按「学校+科室+角色」原文判重；**原文不写库**）" % tot)
    print()
    print("将写入的维度与覆盖率（**缺失不是问题，只要求不被误读**）：")
    for label, key in (("关注方向（兴趣域码）", "interests"),
                       ("学历阶段", "degree"),
                       ("是否临床医学类", "clinical")):
        n = sum(1 for r in recs if r[key])
        print("  %-22s %3d / %3d  (%.0f%%)%s"
              % (label, n, tot, 100.0 * n / tot,
                 "   ← 只在**明确**时映射，映射不了就留空" if key != "interests" else ""))
    n_state = sum(1 for r in recs if r["state"])
    print("  %-22s %3d / %3d  (100%%)" % ("参与状态（新增字段）", n_state, tot))
    print()
    allc = Counter(x for r in recs for x in r["interests"])
    print("关注方向码分布（共 %d 个不同码）：" % len(allc))
    for cc, n in allc.most_common(14):
        print("    %-6s %d 人" % (cc, n))
    unm = Counter(x for r in recs for x in r["unmapped"])
    print()
    print("**未映射**的片段（%d 类）—— 如实计数，不硬塞进任何码：" % len(unm))
    if unm:
        for w, n in unm.most_common(12):
            print("    %-18s %d 人" % (w[:16], n))
    else:
        print("    （无）")
    print()
    print("**不写入**：称呼/姓名、联系方式、学校+科室+角色原文、报名原表任何逐字内容。")
    print("预计新增：%d 个 person 行 + 约 %d 行 field_value"
          % (tot, sum((1 if r["interests"] else 0) + (1 if r["degree"] else 0)
                      + (1 if r["clinical"] else 0) + 1 for r in recs)))


def main():
    ap = argparse.ArgumentParser(description="导入真实报名表（去标识，选项 A）")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--plan", action="store_true")
    g.add_argument("--apply", action="store_true")
    g.add_argument("--cleanup", action="store_true")
    a = ap.parse_args()
    if a.cleanup:
        cleanup()
        return 0
    recs = parse()
    if not recs:
        print("[!] 没解析到任何数据")
        return 1
    if a.plan:
        plan(recs)
        return 0
    plan(recs)
    print()
    apply_(recs)
    return 0


if __name__ == "__main__":
    sys.exit(main())
